#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build báo cáo đồ án RLVD (.docx) theo cấu trúc 7 chương mới.

Khác bản cũ: MỤC LỤC là hyperlink nội bộ thật (bookmark _TocN + w:hyperlink
w:anchor + trường PAGEREF), click được ngay trong Word/LibreOffice mà không
cần "Update Field"; đồng thời bật w:updateFields để số trang tự điền khi mở.

Chạy:  myenv/bin/python build_thesis.py
Kết quả: docs/RLVD_DoAn_2026.docx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).parent
REPO = HERE.parent.parent
sys.path.insert(0, str(HERE))

from thesis_content import BLOCKS, STYLES, PAGE  # noqa: E402

FONT = "Times New Roman"
INK = "1A1A1A"
NAVY = "1F3864"
BLUE = "2E5395"


# ---------------------------------------------------------------- utilities
def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def _add_field(para, instr: str, placeholder: str = "") -> None:
    """Chèn trường Word (BEGIN / instrText / SEPARATE / result / END)."""
    r1 = para.add_run()
    r1._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = para.add_run()
    it = _el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = para.add_run()
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    if placeholder:
        para.add_run(placeholder)
    r4 = para.add_run()
    r4._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def _bookmark(para, name: str, bid: int) -> None:
    """Đóng gói bookmark quanh nội dung paragraph (đầu → cuối)."""
    start = _el("w:bookmarkStart", **{"w:id": str(bid), "w:name": name})
    end = _el("w:bookmarkEnd", **{"w:id": str(bid)})
    para._p.insert(0, start)
    para._p.append(end)


def _internal_hyperlink(para, anchor: str, text_runs) -> None:
    """Thêm w:hyperlink trỏ tới bookmark nội bộ, chứa các run cho sẵn."""
    link = _el("w:hyperlink", **{"w:anchor": anchor, "w:history": "1"})
    for txt, bold in text_runs:
        r = _el("w:r")
        rpr = _el("w:rPr")
        rfonts = _el("w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT,
                                    "w:cs": FONT})
        rpr.append(rfonts)
        sz = _el("w:sz", **{"w:val": "26"})
        rpr.append(sz)
        color = _el("w:color", **{"w:val": "0B3C7A"})
        rpr.append(color)
        if bold:
            rpr.append(_el("w:b"))
        r.append(rpr)
        t = _el("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = txt
        r.append(t)
        link.append(r)
    para._p.append(link)


def _pageref_in_link(para, anchor: str, bid: int) -> None:
    """Trường PAGEREF đặt trong cùng w:hyperlink để số trang cũng click được."""
    link = _el("w:hyperlink", **{"w:anchor": anchor, "w:history": "1"})
    para._p.append(link)

    def _run(children) -> OxmlElement:
        r = _el("w:r")
        rpr = _el("w:rPr")
        rpr.append(_el("w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT}))
        rpr.append(_el("w:sz", **{"w:val": "26"}))
        rpr.append(_el("w:color", **{"w:val": "0B3C7A"}))
        r.append(rpr)
        for c in children:
            r.append(c)
        return r

    link.append(_run([_el("w:fldChar", **{"w:fldCharType": "begin"})]))
    it = _el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" PAGEREF _Toc{bid} \\h "
    link.append(_run([it]))
    link.append(_run([_el("w:fldChar", **{"w:fldCharType": "separate"})]))
    t = _el("w:t")
    t.text = "…"
    link.append(_run([t]))
    link.append(_run([_el("w:fldChar", **{"w:fldCharType": "end"})]))


def _tab_stop_right_dot(para) -> None:
    usable = PAGE["width_mm"] - PAGE["margins_mm"]["left"] - PAGE["margins_mm"]["right"]
    para.paragraph_format.tab_stops.add_tab_stop(
        Mm(usable), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def _enable_update_fields(doc) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        settings.append(_el("w:updateFields", **{"w:val": "true"}))


# ---------------------------------------------------------------- styles
def add_styles(doc) -> None:
    for s in STYLES:
        st = doc.styles.add_style(s["name"], WD_STYLE_TYPE.PARAGRAPH)
        if s.get("base"):
            st.base_style = doc.styles[s["base"]]
        f = st.font
        f.name = s.get("font", FONT)
        if s.get("size_pt"):
            f.size = Pt(s["size_pt"])
        if s.get("bold") is not None:
            f.bold = s["bold"]
        if s.get("italic") is not None:
            f.italic = s["italic"]
        if s.get("color"):
            f.color.rgb = RGBColor.from_string(s["color"])
        pf = st.paragraph_format
        if s.get("space_before_pt") is not None:
            pf.space_before = Pt(s["space_before_pt"])
        if s.get("space_after_pt") is not None:
            pf.space_after = Pt(s["space_after_pt"])
        if s.get("line_spacing"):
            pf.line_spacing = s["line_spacing"]
        if s.get("keep_with_next"):
            pf.keep_with_next = True
        if s.get("justify"):
            pf.alignment = 3  # WD_ALIGN_PARAGRAPH.JUSTIFY
        if s.get("left_indent_mm"):
            pf.left_indent = Mm(s["left_indent_mm"])
        # Font Á-Âu: ép cả eastAsia/cs để Word không rơi về Calibri
        rpr = st.element.get_or_add_rPr()
        rpr.append(_el("w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT,
                                      "w:eastAsia": FONT, "w:cs": FONT}))

    # Normal + danh sách: ép Times New Roman để toàn tài liệu nhất quán
    # (mẫu mặc định của python-docx là Calibri 11 — List Bullet/Number kế thừa nó).
    for nm in ("Normal", "List Bullet", "List Number"):
        try:
            st = doc.styles[nm]
        except KeyError:
            continue
        st.font.name = FONT
        st.font.size = Pt(13)
        st.font.color.rgb = RGBColor.from_string(INK)
        pf = st.paragraph_format
        pf.space_after = Pt(4)
        pf.line_spacing = 1.15
        rpr = st.element.get_or_add_rPr()
        for old in rpr.findall(qn("w:rFonts")):
            rpr.remove(old)
        rpr.append(_el("w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT,
                                      "w:eastAsia": FONT, "w:cs": FONT}))

    # Heading 1-3 chuẩn (màu + font + outline level để TOC \o bắt được)
    heads = {1: (16, True, NAVY, 18, 10), 2: (14, True, BLUE, 14, 6),
             3: (13, True, "333333", 10, 4)}
    for lvl, (size, bold, color, before, after) in heads.items():
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        rpr = st.element.get_or_add_rPr()
        for old in rpr.findall(qn("w:rFonts")):
            rpr.remove(old)
        rpr.append(_el("w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT,
                                      "w:eastAsia": FONT, "w:cs": FONT}))


# ---------------------------------------------------------------- blocks
def add_runs(para, block) -> None:
    runs = block.get("runs")
    if runs is None:
        runs = [{"text": block.get("text", "")}]
    for r in runs:
        run = para.add_run(r.get("text", ""))
        if r.get("bold"):
            run.bold = True
        if r.get("italic"):
            run.italic = True
        if r.get("underline"):
            run.underline = True


def add_table(doc, block) -> None:
    header = block.get("header", [])
    rows = block.get("rows", [])
    ncols = len(header) if header else (len(rows[0]) if rows else 1)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = block.get("style", "Table Grid")
    table.autofit = True
    if header:
        cells = table.add_row().cells
        for i, text in enumerate(header):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(text))
            run.bold = True
            run.font.name = FONT
            run.font.size = Pt(11)
        # tô nền xám nhạt cho hàng tiêu đề
        for c in cells:
            shd = _el("w:shd", **{"w:val": "clear", "w:color": "auto",
                                  "w:fill": "EDF2F8"})
            c._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            if i >= ncols:
                break
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(text))
            run.font.name = FONT
            run.font.size = Pt(11)
    if block.get("caption"):
        cap = doc.add_paragraph(block["caption"], style="CaptionV")
        cap.alignment = 1


def build(out_path: Path) -> dict:
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Mm(PAGE["width_mm"])
    sec.page_height = Mm(PAGE["height_mm"])
    for side in ("top", "bottom", "left", "right"):
        setattr(sec, f"{side}_margin", Mm(PAGE["margins_mm"][side]))
    add_styles(doc)

    # header / footer
    hp = sec.header.paragraphs[0]
    hp.text = ""
    hr = hp.add_run("Đồ án tốt nghiệp — Hệ thống phát hiện vi phạm vượt đèn đỏ (RLVD)")
    hr.font.name = FONT
    hr.font.size = Pt(10)
    hr.font.color.rgb = RGBColor.from_string("666666")
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = 1
    fr = fp.add_run("Trang ")
    fr.font.name = FONT
    fr.font.size = Pt(10)
    _add_field(fp, " PAGE ", "1")
    fr2 = fp.add_run(" / ")
    fr2.font.name = FONT
    fr2.font.size = Pt(10)
    _add_field(fp, " NUMPAGES ", "1")
    for r in fp.runs:
        r.font.name = FONT
        r.font.size = Pt(10)

    # ---- pass 1: đánh số bookmark cho heading (theo thứ tự xuất hiện)
    heading_index: dict[int, tuple[int, str, str]] = {}
    bid = 1000
    for i, b in enumerate(BLOCKS):
        if b.get("type") == "heading":
            heading_index[i] = (bid, b["text"], b.get("level", 1))
            bid += 1

    # ---- pass 2: render
    n_img = n_tab = 0
    for i, b in enumerate(BLOCKS):
        t = b.get("type")
        if t == "heading":
            b_id, text, lvl = heading_index[i]
            p = doc.add_heading(text, level=lvl)
            _bookmark(p, f"_Toc{b_id}", b_id)
        elif t == "paragraph":
            p = doc.add_paragraph(style=b.get("style", "Body"))
            add_runs(p, b)
            if b.get("align") == "center":
                p.alignment = 1
            elif b.get("align") == "justify":
                p.alignment = 3
        elif t == "bullet_list":
            for item in b.get("items", []):
                doc.add_paragraph(item, style="List Bullet")
        elif t == "numbered_list":
            for item in b.get("items", []):
                doc.add_paragraph(item, style="List Number")
        elif t == "plain_list":
            # Danh sách không bullet (danh sách hình/bảng, tài liệu tham khảo)
            for item in b.get("items", []):
                p = doc.add_paragraph(style=b.get("style", "Body"))
                add_runs(p, {"runs": [{"text": item}]})
                pf = p.paragraph_format
                pf.left_indent = Mm(b.get("indent_mm", 8))
                pf.first_line_indent = Mm(-b.get("indent_mm", 8))
                pf.space_after = Pt(2)
        elif t == "table":
            add_table(doc, b)
            n_tab += 1
        elif t == "image":
            path = Path(b["path"])
            if not path.exists():
                raise SystemExit(f"THIẾU ẢNH: {path}")
            doc.add_picture(str(path),
                            width=Mm(b.get("width_mm", 150)))
            doc.paragraphs[-1].alignment = 1
            n_img += 1
            if b.get("caption"):
                doc.add_paragraph(b["caption"], style="CaptionV").alignment = 1
        elif t == "page_break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "toc":
            max_lvl = b.get("max_level", 2)
            for i2, (b_id, text, lvl) in sorted(heading_index.items()):
                if lvl > max_lvl:
                    continue
                p = doc.add_paragraph(style="TocEntry")
                p.paragraph_format.left_indent = Mm(6 * (lvl - 1))
                _tab_stop_right_dot(p)
                _internal_hyperlink(p, f"_Toc{b_id}", [(text, lvl == 1)])
                p.add_run("\t")
                _pageref_in_link(p, f"_Toc{b_id}", b_id)
        else:
            raise ValueError(f"unknown block type: {t}")

    _enable_update_fields(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {"blocks": len(BLOCKS), "headings": len(heading_index),
            "images": n_img, "tables": n_tab,
            "bytes": out_path.stat().st_size, "out": str(out_path)}


if __name__ == "__main__":
    out = REPO / "docs" / "RLVD_DoAn_2026.docx"
    print(json.dumps(build(out), ensure_ascii=False, indent=1))
